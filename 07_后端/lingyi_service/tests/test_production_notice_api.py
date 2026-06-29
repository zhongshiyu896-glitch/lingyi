"""API tests for small-factory production notices."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
from zipfile import ZipFile
from datetime import date
from decimal import Decimal
import os
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image as PILImage
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.main as main_module
from app.main import app
from app.models.audit import Base as AuditBase
from app.models.audit import LyOperationAuditLog
from app.models.bom import Base as BomBase
from app.models.bom import LyFoundationTemplate
from app.models.bom import LyFoundationTemplateNode
from app.models.production import Base as ProductionBase
from app.models.production import LyProductionNotice
from app.models.production import LyProductionPlan
from app.models.production import LyProductionPlanMaterial
from app.models.production import LyProductionPlanOperation
from app.models.sales_order import Base as SalesOrderBase
from app.models.sales_order import LySalesOrder
from app.models.sales_order import LySalesOrderItem
from app.models.style_master import Base as StyleMasterBase
from app.models.style_master import LyStyleGallery
from app.models.style_master import LyStyleMaster
from app.routers.auth import get_db_session as auth_db_dep
from app.routers.production import get_db_session as production_db_dep


class ProductionNoticeApiTest(unittest.TestCase):
    """Validate production notice creation is order-scoped and non-mutating."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.engine = create_engine(
            "sqlite+pysqlite://",
            future=True,
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
            execution_options={"schema_translate_map": {"ly_schema": None, "public": None}},
        )
        cls.SessionLocal = sessionmaker(bind=cls.engine, autoflush=False, autocommit=False, expire_on_commit=False)
        BomBase.metadata.create_all(bind=cls.engine)
        StyleMasterBase.metadata.create_all(bind=cls.engine)
        SalesOrderBase.metadata.create_all(bind=cls.engine)
        ProductionBase.metadata.create_all(bind=cls.engine)
        AuditBase.metadata.create_all(bind=cls.engine)

        def _override_db():
            db = cls.SessionLocal()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[auth_db_dep] = _override_db
        app.dependency_overrides[production_db_dep] = _override_db
        cls._old_main_session_local = main_module.SessionLocal
        cls._old_upload_dir = os.environ.get("LINGYI_UPLOAD_DIR")
        main_module.SessionLocal = cls.SessionLocal
        cls.client = TestClient(app)
        cls.upload_dir = tempfile.TemporaryDirectory()

    @classmethod
    def tearDownClass(cls) -> None:
        if cls._old_upload_dir is None:
            os.environ.pop("LINGYI_UPLOAD_DIR", None)
        else:
            os.environ["LINGYI_UPLOAD_DIR"] = cls._old_upload_dir
        main_module.SessionLocal = cls._old_main_session_local
        app.dependency_overrides.pop(auth_db_dep, None)
        app.dependency_overrides.pop(production_db_dep, None)
        cls.engine.dispose()
        cls.upload_dir.cleanup()

    def setUp(self) -> None:
        os.environ["APP_ENV"] = "test"
        os.environ["LINGYI_ALLOW_DEV_AUTH"] = "true"
        os.environ["LINGYI_PERMISSION_SOURCE"] = "static"
        os.environ["LINGYI_ERPNEXT_BASE_URL"] = ""
        os.environ["LINGYI_UPLOAD_DIR"] = self.upload_dir.name
        with self.SessionLocal() as session:
            session.query(LyOperationAuditLog).delete()
            session.query(LyProductionPlanOperation).delete()
            session.query(LyProductionPlanMaterial).delete()
            session.query(LyProductionPlan).delete()
            session.query(LyProductionNotice).delete()
            session.query(LySalesOrderItem).delete()
            session.query(LySalesOrder).delete()
            session.query(LyFoundationTemplateNode).delete()
            session.query(LyFoundationTemplate).delete()
            session.query(LyStyleGallery).delete()
            session.query(LyStyleMaster).delete()
            session.commit()
        self._write_upload_image("images/style_gallery/test-style.png", color=(222, 184, 135))
        self._write_upload_image("images/style_gallery/test-wash.png", color=(245, 245, 245))
        self._seed_order_and_templates()

    @staticmethod
    def _headers(request_id: str = "PROD-NOTICE-REQ") -> dict[str, str]:
        return {
            "X-LY-Dev-User": "prod.notice.user",
            "X-LY-Dev-Roles": "Production Manager",
            "X-Request-ID": request_id,
        }

    @staticmethod
    def _xlsx_sheet_xml(content: bytes) -> str:
        with ZipFile(BytesIO(content)) as workbook:
            return workbook.read("xl/worksheets/sheet1.xml").decode("utf-8")

    @staticmethod
    def _xlsx_workbook(content: bytes):
        return load_workbook(BytesIO(content))

    def _write_upload_image(self, relative_path: str, *, color: tuple[int, int, int]) -> None:
        path = Path(self.upload_dir.name) / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        image = PILImage.new("RGB", (160, 120), color=color)
        image.save(path, format="PNG")

    def _seed_order_and_templates(self) -> None:
        with self.SessionLocal() as session:
            session.add(
                LyStyleMaster(
                    id=7100,
                    company="COMP-N",
                    ys_style_no="STYLE-NOTICE-001",
                    ys_style_name_cn="通知单款",
                    ys_season="夏",
                    ys_year="2026",
                    ys_brand="LY",
                    ys_style_status="enabled",
                    colors=["白", "黑"],
                    sizes=["S", "M"],
                    size_chart={
                        "unit": "CM",
                        "sizes": ["S", "M"],
                        "rows": [
                            {"part": "后中", "values": {"S": "62.1", "M": "64.1"}, "sort_no": 10},
                            {"part": "肩宽", "values": {"S": "48", "M": "49.2"}, "sort_no": 20},
                        ],
                    },
                    created_by="seed",
                    updated_by="seed",
                )
            )
            session.add(
                LyStyleGallery(
                    id=7101,
                    company="COMP-N",
                    style_master_id=7100,
                    image_url="/uploads/images/style_gallery/test-style.png",
                    image_name="主图",
                    image_type="main",
                    is_primary=True,
                    status="active",
                    created_by="seed",
                )
            )
            session.add(
                LyStyleGallery(
                    id=7102,
                    company="COMP-N",
                    style_master_id=7100,
                    image_url="/uploads/images/style_gallery/test-wash.png",
                    image_name="水洗标",
                    image_type="wash_label",
                    is_primary=False,
                    status="active",
                    created_by="seed",
                )
            )
            order = LySalesOrder(
                id=7200,
                sales_order_no="SO-NOTICE-001",
                source_order_ref="SO-NOTICE-001",
                company="COMP-N",
                customer="通知单客户",
                status="planned",
                docstatus=1,
                transaction_date=date(2026, 6, 27),
                delivery_date=date(2026, 7, 7),
                currency="CNY",
                grand_total=Decimal("0"),
                idempotency_key="idem-notice-order",
                request_hash="hash-notice-order",
                created_by="seed",
            )
            session.add(order)
            for line_id, line_no, color, size, qty in [
                (7201, 1, "白", "S", Decimal("100")),
                (7202, 2, "黑", "M", Decimal("200")),
            ]:
                session.add(
                    LySalesOrderItem(
                        id=line_id,
                        sales_order_id=7200,
                        company="COMP-N",
                        line_no=line_no,
                        sales_order_item=f"SO-NOTICE-001-{line_no:03d}",
                        style_master_id=7100,
                        item_code="STYLE-NOTICE-001",
                        item_name="通知单款",
                        color=color,
                        size=size,
                        qty=qty,
                        planned_qty=qty,
                        delivered_qty=Decimal("0"),
                        ys_material_calc_state="待算料",
                        uom="Nos",
                        delivery_date=date(2026, 7, 7),
                    )
                )
            workmanship = LyFoundationTemplate(
                id=7300,
                company="COMP-N",
                template_type="workmanship",
                template_code="PROC-001",
                name="基础工艺",
                scene="通用",
                status="active",
                created_by="seed",
            )
            size_chart = LyFoundationTemplate(
                id=7310,
                company="COMP-N",
                template_type="size_spec",
                template_code="SIZE-001",
                name="基础尺寸表",
                scene="通用",
                status="active",
                created_by="seed",
            )
            session.add_all([workmanship, size_chart])
            session.add_all(
                [
                    LyFoundationTemplateNode(
                        id=7301,
                        template_id=7300,
                        code="PROC-ROW-1",
                        name="车线平整",
                        node_type="工艺说明",
                        owner="业务",
                        sort_no=10,
                        status="active",
                        created_by="seed",
                    ),
                    LyFoundationTemplateNode(
                        id=7311,
                        template_id=7310,
                        code="SIZE-ROW-1",
                        name="胸围",
                        node_type="尺寸项目",
                        owner="CM",
                        sort_no=10,
                        status="active",
                        created_by="seed",
                    ),
                ]
            )
            session.commit()

    def _seed_ready_plan(
        self,
        *,
        plan_id: int = 7400,
        plan_no: str = "PP-NOTICE-001",
        sales_order_item: str = "SO-NOTICE-001-001",
    ) -> int:
        with self.SessionLocal() as session:
            session.add(
                LyProductionPlan(
                    id=plan_id,
                    plan_no=plan_no,
                    plan_group_no="PPG-NOTICE-001",
                    company="COMP-N",
                    sales_order="SO-NOTICE-001",
                    sales_order_item=sales_order_item,
                    customer="通知单客户",
                    item_code="STYLE-NOTICE-001",
                    bom_id=7100,
                    bom_version="V1",
                    planned_qty=Decimal("100"),
                    planned_start_date=date(2026, 6, 28),
                    status="material_checked",
                    idempotency_key=f"idem-{plan_no}",
                    request_hash=f"hash-{plan_no}",
                    created_by="seed",
                )
            )
            session.flush()
            session.add(
                LyProductionPlanMaterial(
                    plan_id=plan_id,
                    bom_item_id=1,
                    bom_color="白",
                    bom_size="S",
                    bom_part="主面",
                    material_item_code="FAB-NOTICE",
                    warehouse="WH-N",
                    uom="米",
                    qty_per_piece=Decimal("1"),
                    loss_rate=Decimal("0"),
                    required_qty=Decimal("100"),
                    available_qty=Decimal("100"),
                    shortage_qty=Decimal("0"),
                )
            )
            session.commit()
        return plan_id

    def _start_production_payload(
        self,
        *,
        plan_id: int,
        idempotency_key: str,
        sales_order_item: str = "SO-NOTICE-001-001",
    ) -> dict[str, object]:
        return {
            "company": "COMP-N",
            "action": "start",
            "remark": "齐料后开始生产",
            "production_mode": "in_house",
            "factory_name": "本厂",
            "production_start_date": "2026-06-28",
            "operation": "production_status",
            "scenario_tag": "production_status",
            "idempotency_key": idempotency_key,
            "plan_id": plan_id,
            "sales_order": "SO-NOTICE-001",
            "sales_order_item": sales_order_item,
            "item_code": "STYLE-NOTICE-001",
        }

    def _post_start_production(self, *, plan_id: int, idempotency_key: str, request_id: str, sales_order_item: str = "SO-NOTICE-001-001"):
        return self.client.post(
            f"/api/production/plans/{plan_id}/production-status",
            json=self._start_production_payload(
                plan_id=plan_id,
                idempotency_key=idempotency_key,
                sales_order_item=sales_order_item,
            ),
            headers=self._headers(request_id),
        )

    def _assert_notice_gate_blocked(self, response) -> None:
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(response.json()["code"], "PRODUCTION_TRACKING_NODE_INVALID")
        self.assertEqual(response.json()["message"], "请先确认/下发生产通知单，再开始生产")

    def test_notice_create_update_export_and_order_state_is_not_mutated(self) -> None:
        payload = {
            "company": "COMP-N",
            "sales_order": "SO-NOTICE-001",
            "factory_name": "默认加工厂",
            "workmanship_template_id": 7300,
            "size_template_id": 7310,
            "process_text": "按样衣生产",
            "packaging_text": "单件入袋",
            "label_text": "按唛头要求",
        }
        response = self.client.post("/api/production/notices", json=payload, headers=self._headers())
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["code"], "0")
        data = body["data"]
        self.assertEqual(data["sales_order"], "SO-NOTICE-001")
        self.assertEqual(data["order_qty"], "300.000000")
        self.assertEqual(len(data["color_size_matrix"]), 2)
        self.assertEqual(data["style_image_url"], "/uploads/images/style_gallery/test-style.png")
        self.assertEqual(data["workmanship_snapshot"]["nodes"][0]["name"], "车线平整")

        repeat_response = self.client.post(
            "/api/production/notices",
            json={**payload, "factory_name": "改后的加工厂"},
            headers=self._headers("PROD-NOTICE-REQ-REPEAT"),
        )
        self.assertEqual(repeat_response.status_code, 200)
        repeat_data = repeat_response.json()["data"]
        self.assertEqual(repeat_data["id"], data["id"])
        self.assertEqual(repeat_data["factory_name"], "改后的加工厂")

        patch_response = self.client.patch(
            f"/api/production/notices/{data['id']}",
            json={
                "company": "COMP-N",
                "factory_name": "二次编辑加工厂",
                "process_text": "二次编辑工艺",
                "packaging_text": "二次编辑包装",
                "label_text": "二次编辑唛头",
                "remark": "二次编辑备注",
                "status": "sent",
            },
            headers=self._headers("PROD-NOTICE-REQ-PATCH"),
        )
        self.assertEqual(patch_response.status_code, 200)
        patched_data = patch_response.json()["data"]
        self.assertEqual(patched_data["status"], "sent")
        self.assertEqual(patched_data["factory_name"], "二次编辑加工厂")
        self.assertEqual(patched_data["process_text"], "二次编辑工艺")

        list_response = self.client.get(
            "/api/production/notices",
            params={"company": "COMP-N", "keyword": "SO-NOTICE-001"},
            headers=self._headers("PROD-NOTICE-REQ-LIST"),
        )
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(list_response.json()["data"]["total"], 1)

        export_response = self.client.get(
            f"/api/production/notices/{data['id']}/export.xlsx",
            params={"company": "COMP-N"},
            headers=self._headers("PROD-NOTICE-REQ-EXPORT"),
        )
        self.assertEqual(export_response.status_code, 200)
        self.assertTrue(export_response.content.startswith(b"PK"))
        sheet_xml = self._xlsx_sheet_xml(export_response.content)
        self.assertIn("二次编辑加工厂", sheet_xml)
        self.assertIn("二次编辑工艺", sheet_xml)
        self.assertIn("二次编辑包装", sheet_xml)
        self.assertIn("二次编辑唛头", sheet_xml)
        self.assertIn("二次编辑备注", sheet_xml)
        self.assertNotIn("默认加工厂", sheet_xml)
        self.assertNotIn("按样衣生产", sheet_xml)
        workbook = self._xlsx_workbook(export_response.content)
        worksheet = workbook.active
        self.assertEqual(workbook.sheetnames, ["Sheet1"])
        self.assertEqual(worksheet["A1"].value, "杭州领意服饰有限公司生产通知单（ STYLE-NOTICE-001 ）")
        self.assertEqual(worksheet["H2"].value, "版号：STYLE-NOTICE-001")
        self.assertEqual(worksheet["M3"].value, "2026年7月7日")
        self.assertEqual(worksheet["C17"].value, "100")
        self.assertEqual(worksheet["D17"].value, "200")
        self.assertEqual(worksheet["C19"].value, "100")
        self.assertEqual(worksheet["D19"].value, "200")
        self.assertEqual(worksheet["A22"].value, "后中")
        self.assertEqual(worksheet["C22"].value, "62.1")
        self.assertEqual(str(worksheet.print_area), "'Sheet1'!$A$1:$V$28")
        self.assertGreaterEqual(len(worksheet.merged_cells.ranges), 46)
        self.assertGreaterEqual(len(worksheet._images), 2)

        with self.SessionLocal() as session:
            self.assertEqual(session.query(LyProductionNotice).count(), 1)
            order = session.query(LySalesOrder).filter_by(sales_order_no="SO-NOTICE-001").one()
            self.assertEqual(order.status, "planned")
            self.assertEqual(order.grand_total, Decimal("0.000000"))

    def test_notice_template_export_without_images_and_missing_fields_keeps_blanks(self) -> None:
        with self.SessionLocal() as session:
            session.query(LyStyleGallery).delete()
            style = session.query(LyStyleMaster).filter_by(id=7100).one()
            style.size_chart = {}
            session.commit()

        create_response = self.client.post(
            "/api/production/notices",
            json={"company": "COMP-N", "sales_order": "SO-NOTICE-001"},
            headers=self._headers("PROD-NOTICE-EXPORT-NO-IMAGE-CREATE"),
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        notice_id = create_response.json()["data"]["id"]

        export_response = self.client.get(
            f"/api/production/notices/{notice_id}/export.xlsx",
            params={"company": "COMP-N"},
            headers=self._headers("PROD-NOTICE-EXPORT-NO-IMAGE"),
        )
        self.assertEqual(export_response.status_code, 200, export_response.text)
        workbook = self._xlsx_workbook(export_response.content)
        worksheet = workbook.active
        self.assertEqual(worksheet["A1"].value, "杭州领意服饰有限公司生产通知单（ STYLE-NOTICE-001 ）")
        self.assertIsNone(worksheet["D7"].value)
        self.assertIsNone(worksheet["A22"].value)
        self.assertEqual(str(worksheet.print_area), "'Sheet1'!$A$1:$V$28")
        self.assertEqual(len(worksheet._images), 0)
        self.assertGreaterEqual(len(worksheet.merged_cells.ranges), 46)

    def test_notice_create_defaults_factory_to_in_house_and_repeat_does_not_unbind(self) -> None:
        first = self.client.post(
            "/api/production/notices",
            json={"company": "COMP-N", "sales_order": "SO-NOTICE-001"},
            headers=self._headers("PROD-NOTICE-DEFAULT-FACTORY"),
        )
        self.assertEqual(first.status_code, 200, first.text)
        first_data = first.json()["data"]
        self.assertEqual(first_data["factory_name"], "本厂")

        repeat = self.client.post(
            "/api/production/notices",
            json={"company": "COMP-N", "sales_order": "SO-NOTICE-001"},
            headers=self._headers("PROD-NOTICE-DEFAULT-FACTORY-REPEAT"),
        )
        self.assertEqual(repeat.status_code, 200, repeat.text)
        repeat_data = repeat.json()["data"]
        self.assertEqual(repeat_data["id"], first_data["id"])
        self.assertEqual(repeat_data["factory_name"], "本厂")

    def test_notice_status_gate_controls_production_start(self) -> None:
        plan_id = self._seed_ready_plan()

        no_notice = self._post_start_production(
            plan_id=plan_id,
            idempotency_key="idem-notice-start-none",
            request_id="PROD-NOTICE-START-NONE",
        )
        self._assert_notice_gate_blocked(no_notice)

        notice = self.client.post(
            "/api/production/notices",
            json={"company": "COMP-N", "sales_order": "SO-NOTICE-001", "status": "draft"},
            headers=self._headers("PROD-NOTICE-CREATE-DRAFT"),
        )
        self.assertEqual(notice.status_code, 200, notice.text)
        self.assertEqual(notice.json()["data"]["status"], "draft")

        draft_notice = self._post_start_production(
            plan_id=plan_id,
            idempotency_key="idem-notice-start-draft",
            request_id="PROD-NOTICE-START-DRAFT",
        )
        self._assert_notice_gate_blocked(draft_notice)

        confirm = self.client.patch(
            f"/api/production/notices/{notice.json()['data']['id']}",
            json={"company": "COMP-N", "status": "confirmed"},
            headers=self._headers("PROD-NOTICE-CONFIRM"),
        )
        self.assertEqual(confirm.status_code, 200, confirm.text)
        self.assertEqual(confirm.json()["data"]["status"], "confirmed")

        confirmed_start = self._post_start_production(
            plan_id=plan_id,
            idempotency_key="idem-notice-start-confirmed",
            request_id="PROD-NOTICE-START-CONFIRMED",
        )
        self.assertEqual(confirmed_start.status_code, 200, confirmed_start.text)
        self.assertEqual(confirmed_start.json()["data"]["status"], "production_in_progress")
        self.assertEqual(confirmed_start.json()["data"]["production_notice"]["notice_status"], "confirmed")

        second_plan_id = self._seed_ready_plan(
            plan_id=7401,
            plan_no="PP-NOTICE-002",
            sales_order_item="SO-NOTICE-001-002",
        )
        sent = self.client.patch(
            f"/api/production/notices/{notice.json()['data']['id']}",
            json={"company": "COMP-N", "status": "sent"},
            headers=self._headers("PROD-NOTICE-SEND"),
        )
        self.assertEqual(sent.status_code, 200, sent.text)
        self.assertEqual(sent.json()["data"]["status"], "sent")

        sent_start = self._post_start_production(
            plan_id=second_plan_id,
            idempotency_key="idem-notice-start-sent",
            request_id="PROD-NOTICE-START-SENT",
            sales_order_item="SO-NOTICE-001-002",
        )
        self.assertEqual(sent_start.status_code, 200, sent_start.text)
        self.assertEqual(sent_start.json()["data"]["status"], "production_in_progress")
        self.assertEqual(sent_start.json()["data"]["production_notice"]["notice_status"], "sent")

    def test_repeat_create_does_not_downgrade_confirmed_notice_without_explicit_status(self) -> None:
        first = self.client.post(
            "/api/production/notices",
            json={"company": "COMP-N", "sales_order": "SO-NOTICE-001", "status": "confirmed"},
            headers=self._headers("PROD-NOTICE-CREATE-CONFIRMED"),
        )
        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(first.json()["data"]["status"], "confirmed")

        repeat_without_status = self.client.post(
            "/api/production/notices",
            json={"company": "COMP-N", "sales_order": "SO-NOTICE-001", "factory_name": "二次更新工厂"},
            headers=self._headers("PROD-NOTICE-REPEAT-NO-STATUS"),
        )
        self.assertEqual(repeat_without_status.status_code, 200, repeat_without_status.text)
        self.assertEqual(repeat_without_status.json()["data"]["id"], first.json()["data"]["id"])
        self.assertEqual(repeat_without_status.json()["data"]["status"], "confirmed")
        self.assertEqual(repeat_without_status.json()["data"]["factory_name"], "二次更新工厂")


if __name__ == "__main__":
    unittest.main()
